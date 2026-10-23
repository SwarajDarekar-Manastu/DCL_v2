"""A1.9–A1.10: Word files via LibreOffice and the PDF pipeline plus DOCX extras; slides read natively."""

import io
import zipfile

import pytest

pikepdf = pytest.importorskip("pikepdf")
pptx = pytest.importorskip("pptx")

from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from backend.postgres.enums import AssetKind, NodeSource, NodeType  # noqa: E402
from backend.routes.a1_office.office import (  # noqa: E402
    UnsupportedOfficeFormat,
    apply_word_extras,
    process_office,
    word_extras,
)
from backend.routes.a1_office.route import Stores  # noqa: E402

from .test_route import SHA, fixture, folder, stores, triage  # noqa: E402

PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00"
    b"\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0\x00\x00\x03\x01\x01\x00\xc9\xfe\x92\xef\x00\x00\x00\x00IEND\xaeB`\x82"
)


def blank_pdf(pages: int) -> bytes:
    pdf = pikepdf.new()
    for _ in range(pages):
        pdf.add_blank_page(page_size=(720, 405))
    out = io.BytesIO()
    pdf.save(out)
    return out.getvalue()


def deck(slides: int = 10) -> bytes:
    prs = pptx.Presentation()
    for n in range(1, slides + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Heater test {n}"
        slide.placeholders[1].text_frame.text = f"Preheat to {n * 10} W"
        if n == 2:
            data = CategoryChartData()
            data.categories = ["T01", "T02"]
            data.add_series("Peak Pc (bar)", (12.4, 11.9))
            slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(8), Inches(5), data
            )
        if n == 3:
            slide.shapes.add_picture(io.BytesIO(PNG), Inches(0), Inches(0), prs.slide_width, prs.slide_height)
            slide.notes_slide.notes_text_frame.text = "Mesh burned at the throat."
        if n == 4:
            table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
            for (r, c), text in {(0, 0): "Run", (0, 1): "Isp", (1, 0): "T01", (1, 1): "221 s"}.items():
                table.cell(r, c).text = text
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def converter(calls):
    def convert(data, ext, to):
        calls.append((ext, to))
        if to == "pdf":
            return blank_pdf(10) if ext in ("pptx", "ppt") else fixture("german_datasheet")[1]
        return data

    return convert


DESCRIBED = {"summary": "Photo of the burned injector mesh.", "diagram_description": None, "key_numbers": []}


def test_a_deck_becomes_one_slide_node_per_slide_with_notes_charts_tables_and_renders(
    fake_vision, memory_store
):
    vision = fake_vision(triage_page=triage(), describe_slide=DESCRIBED)
    calls = []
    doc = process_office(
        deck(),
        blob_sha256=SHA,
        file_name="Heater review.pptx",
        folder=folder("APU Heater", "Workshops", "Heater review.pptx"),
        vision=vision,
        stores=stores(memory_store),
        convert_office=converter(calls),
    )
    slides = [n for n in doc.nodes if n.node_type is NodeType.SLIDE]
    assert [s.page_no for s in slides] == list(range(1, 11))
    assert "Peak Pc (bar): T01 = 12.4, T02 = 11.9" in slides[1].content
    assert "Speaker notes: Mesh burned at the throat." in slides[2].content
    assert "Run | Isp\nT01 | 221 s" in slides[3].content
    parents = [n for n in doc.nodes if n.node_type is NodeType.PARENT]
    assert [(p.page_no, p.page_end) for p in parents] == [(1, 8), (9, 10)]
    assert all(s.parent_ref in {p.ref for p in parents} for s in slides)
    [figure] = [n for n in doc.nodes if n.node_type is NodeType.FIGURE]
    assert (
        figure.page_no == 3 and figure.source is NodeSource.VLM and "burned injector mesh" in figure.content
    )
    assert [c for c, _ in vision.calls] == [
        "triage_page",
        "describe_slide",
    ]  # the chart slide is not described
    assert sum(a.kind is AssetKind.PAGE_RENDER for a in doc.assets) == 10
    assert doc.document.page_count == 10 and doc.document.language == "en"
    assert calls == [("pptx", "pdf")]


def test_without_cloud_no_slide_goes_to_gemini_and_legacy_ppt_is_converted_first(fake_vision, memory_store):
    vision = fake_vision()
    calls = []
    doc = process_office(
        deck(),
        blob_sha256=SHA,
        file_name="old.ppt",
        folder=folder("APU", "Workshops", "old.ppt"),
        vision=vision,
        stores=stores(memory_store),
        convert_office=converter(calls),
        cloud_allowed=False,
    )
    assert vision.calls == [] and calls == [("ppt", "pptx"), ("pptx", "pdf")]
    assert not any(n.node_type is NodeType.FIGURE for n in doc.nodes)


W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def word_file(*, tracked=True) -> bytes:
    change = '<w:ins w:id="9" w:author="A"><w:r><w:t> rev B</w:t></w:r></w:ins>' if tracked else ""
    document = (
        f"<w:document {W_NS}><w:body>"
        '<w:p><w:commentRangeStart w:id="1"/><w:r><w:t>Nennschub 10 N bei 20 bar</w:t></w:r>'
        f'<w:commentRangeEnd w:id="1"/>{change}<w:r><w:footnoteReference w:id="2"/></w:r></w:p>'
        "</w:body></w:document>"
    )
    comments = (
        f'<w:comments {W_NS}><w:comment w:id="1" w:author="Swaraj">'
        "<w:p><w:r><w:t>Check against TP7</w:t></w:r></w:p></w:comment></w:comments>"
    )
    footnotes = (
        f'<w:footnotes {W_NS}><w:footnote w:id="2">'
        "<w:p><w:r><w:t>Measured at 20 °C.</w:t></w:r></w:p></w:footnote></w:footnotes>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/comments.xml", comments)
        archive.writestr("word/footnotes.xml", footnotes)
        archive.writestr("word/embeddings/Microsoft_Excel_Worksheet.xlsx", b"x" * 1234)
    return out.getvalue()


def test_word_extras_read_what_only_the_docx_has():
    extras = word_extras(word_file())
    assert extras.has_tracked_changes
    [comment] = extras.comments
    assert (comment.author, comment.text, comment.anchor_text) == (
        "Swaraj",
        "Check against TP7",
        "Nennschub 10 N bei 20 bar rev B",
    )
    assert extras.footnotes == [("Nennschub 10 N bei 20 bar rev B", "Measured at 20 °C.")]
    assert extras.embedded == [{"name": "Microsoft_Excel_Worksheet.xlsx", "bytes": 1234}]
    assert not word_extras(word_file(tracked=False)).has_tracked_changes


def test_a_word_file_runs_the_pdf_pipeline_and_keeps_its_comments_and_footnotes(fake_vision, memory_store):
    calls = []
    doc = process_office(
        word_file(),
        blob_sha256=SHA,
        file_name="Datenblatt.docx",
        folder=folder("APS - Library", "Vendor brochures/DE", "Datenblatt.docx"),
        vision=fake_vision(triage_page=triage(language="de")),
        stores=stores(memory_store),
        convert_office=converter(calls),
    )
    assert calls == [("docx", "pdf")]
    assert any(n.node_type is NodeType.TABLE for n in doc.nodes)  # from the PDF pipeline
    assert doc.document.catalog["word"] == {
        "has_tracked_changes": True,
        "comments": 1,
        "footnotes": 1,
        "embedded_objects": [{"name": "Microsoft_Excel_Worksheet.xlsx", "bytes": 1234}],
    }
    [annotation] = [m for m in doc.mentions if m.kind == "annotation"]
    assert annotation.surface == "Swaraj: Check against TP7"


def test_extras_attach_to_the_node_that_holds_the_paragraph(fake_vision, memory_store):
    from backend.dispatcher.contract import DocumentFields, ProcessedDocument, ProcessedNode

    node = ProcessedNode(
        ref="n1",
        node_type=NodeType.CHILD,
        ordinal=0,
        content="Nennschub  10 N bei 20 bar rev B. More.",
        page_no=1,
        source=NodeSource.PRINTED,
    )
    base = ProcessedDocument(
        blob_sha256=SHA,
        route="A1",
        parser_version=1,
        cir_uri=f"s3://dcl-cir/{SHA}/A1/v1/",
        document=DocumentFields(title="x"),
        nodes=[node],
    )
    out = apply_word_extras(base, word_extras(word_file()))
    assert out.nodes[0].content.endswith("[Footnote: Measured at 20 °C.]")
    assert out.mentions[0].node_ref == "n1"


def test_other_formats_are_refused(fake_vision, memory_store):
    with pytest.raises(UnsupportedOfficeFormat):
        process_office(
            b"x",
            blob_sha256=SHA,
            file_name="a.xyz",
            folder=None,
            vision=fake_vision(),
            stores=Stores(put_cir=lambda k, d: k, put_asset=lambda k, d: k),
            convert_office=lambda *a: b"",
        )

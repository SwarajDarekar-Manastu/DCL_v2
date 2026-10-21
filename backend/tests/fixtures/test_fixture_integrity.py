"""Integrity of the fixture corpus (feature quality/fixtures).

Proves, without running the pipeline, that every fixture has a valid expected.yaml and really
has the property that makes it an edge case. Hostile fixtures are inspected safely: the pickle
only through pickletools, archives only through their central directory.
"""

import ast
import csv
import hashlib
import io
import pickletools
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import jsonschema
import pytest
import yaml
from openpyxl import load_workbook
from pypdf import PdfReader

FIXTURES = Path(__file__).resolve().parent
EXPECTED_FILES = sorted(FIXTURES.glob("*/*/expected.yaml"))

ROUTES = ["A1", "A2", "B", "C", "D", "E", "F", "G", None]
STATUSES = ["ACTIVE", "ACTIVE_WITH_SOFT_FLAGS", "HALTED", "WAITING_FOR_ROUTE", "QUARANTINED"]
FLAG_TYPES = [
    "UNIT_AMBIGUITY",
    "SCHEMA_UNKNOWN",
    "LOW_CONFIDENCE_EXTRACTION",
    "PARSE_FAILURE",
    "UNSUPPORTED_FORMAT",
    "CONFLICT",
    "SENSITIVE_CONTENT",
    "HELD_LARGE_FILE",
    "CLOUD_RESTRICTED",
]
ROOTS = ["APU Injector", "APU", "APS - Library", "APU Heater"]
ROUTE_DIRS = [
    "dispatcher",
    "a1_office",
    "a2_source",
    "b_code",
    "c_telemetry",
    "d_photos",
    "e_diagrams",
    "f_video",
    "g_archives",
]

SCHEMA = {
    "type": "object",
    "required": [
        "fixture",
        "edge_case",
        "origin",
        "byte_stable",
        "files",
        "simulated_location",
        "spec",
        "expected",
    ],
    "additionalProperties": False,
    "properties": {
        "fixture": {"type": "string"},
        "edge_case": {"type": "string"},
        "origin": {"enum": ["synthetic", "real"]},
        "generator": {"type": "string", "pattern": r"^make_synthetic\.py::[a-z_]+$"},
        "byte_stable": {"type": "boolean"},
        "files": {"type": "array", "minItems": 1, "items": {"type": "string"}},
        "simulated_location": {
            "type": "object",
            "required": ["root", "folder"],
            "additionalProperties": False,
            "properties": {"root": {"enum": ROOTS}, "folder": {"type": "string"}},
        },
        "spec": {
            "type": "array",
            "minItems": 1,
            "items": {"type": "string", "pattern": "^docs/architecture/"},
        },
        "notes": {"type": "string"},
        "open_questions": {"type": "array", "items": {"type": "string"}},
        "expected": {
            "type": "object",
            "required": ["route", "detected_mime", "document_status", "flags", "flags_exact"],
            "additionalProperties": False,
            "properties": {
                "route": {"enum": ROUTES},
                "detected_mime": {"type": "string", "pattern": r"^[a-z]+/[a-z0-9.+-]+$"},
                "documents": {"type": "integer", "minimum": 1},
                "document_status": {"enum": STATUSES},
                "flags": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "required": ["type", "severity", "reason"],
                        "additionalProperties": False,
                        "properties": {
                            "type": {"enum": FLAG_TYPES},
                            "severity": {"enum": ["HARD", "SOFT"]},
                            "reason": {"type": "string"},
                        },
                    },
                },
                "flags_exact": {"type": "boolean"},
                "metadata": {"type": "object"},
                "text_contains": {"type": "array", "items": {"type": "string"}},
                "ocr_text_contains": {"type": "array", "items": {"type": "string"}},
                "numeric_facts": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "required": ["quantity", "value", "unit"],
                        "additionalProperties": False,
                        "properties": {
                            "quantity": {"type": "string"},
                            "value": {"type": "number"},
                            "unit": {"type": "string"},
                            "source": {"enum": ["vendor_claim", "literature"]},
                        },
                    },
                },
                "entities": {"type": "array", "items": {"type": "object", "required": ["type", "surface"]}},
                "never": {"type": "array", "items": {"type": "string"}},
            },
        },
    },
}


def load(path: Path) -> dict:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def fixture(name: str) -> tuple[Path, dict]:
    path = next(p for p in EXPECTED_FILES if p.parent.name == name)
    return path.parent, load(path)


def pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


# ---- every fixture ----------------------------------------------------------------------


def test_fixture_set_covers_file_10_section_7():
    names = {p.parent.name for p in EXPECTED_FILES}
    assert {
        "rotated_scan",
        "merged_cells_fmea",
        "headerless_csv",
        "two_latex_masters",
        "zip_bomb_ratio",
        "malicious_pickle",
        "pdf_is_html",
        "german_datasheet",
        "japanese_scanned_page",
        "doi_paper",
    } <= names


@pytest.mark.parametrize("path", EXPECTED_FILES, ids=lambda p: p.parent.name)
def test_expected_yaml_is_valid_and_files_match(path):
    data = load(path)
    jsonschema.validate(data, SCHEMA)
    assert data["fixture"] == path.parent.name
    assert path.parent.parent.name in ROUTE_DIRS
    present = sorted(
        str(p.relative_to(path.parent))
        for p in path.parent.rglob("*")
        if p.is_file() and p.name != "expected.yaml"
    )
    assert present == sorted(data["files"]), "fixture folder must hold exactly the listed files"
    for flag in data["expected"]["flags"]:
        if flag["severity"] == "HARD":
            assert data["expected"]["document_status"] == "HALTED"


@pytest.mark.parametrize(
    "path",
    [p for p in EXPECTED_FILES if load(p)["expected"].get("text_contains")],
    ids=lambda p: p.parent.name,
)
def test_text_contains_is_in_the_file(path):
    data = load(path)
    texts = []
    for name in data["files"]:
        f = path.parent / name
        texts.append(pdf_text(f) if f.suffix == ".pdf" else f.read_text(encoding="utf-8"))
    joined = "\n".join(texts)
    for s in data["expected"]["text_contains"]:
        assert s in joined, f"{s!r} not found in {path.parent.name}"


def test_regeneration_reproduces_byte_stable_fixtures(tmp_path):
    subprocess.run([sys.executable, str(FIXTURES / "make_synthetic.py"), "--out", str(tmp_path)], check=True)
    for path in EXPECTED_FILES:
        data = load(path)
        if data["origin"] != "synthetic":
            continue
        rel = path.parent.relative_to(FIXTURES)
        for name in data["files"]:
            fresh = tmp_path / rel / name
            assert fresh.exists(), f"generator no longer produces {rel / name}"
            if data["byte_stable"]:
                committed = hashlib.sha256((path.parent / name).read_bytes()).hexdigest()
                assert hashlib.sha256(fresh.read_bytes()).hexdigest() == committed, f"{rel / name} drifted"


def test_fixture_corpus_is_small():
    total = sum(p.stat().st_size for p in FIXTURES.rglob("*") if p.is_file())
    assert total < 1_000_000


# ---- per edge case ----------------------------------------------------------------------


@pytest.mark.parametrize("name", ["rotated_scan", "japanese_scanned_page"])
def test_scanned_pages_have_no_text_layer(name):
    folder, data = fixture(name)
    reader = PdfReader(folder / data["files"][0])
    assert len(reader.pages) == 1
    assert (reader.pages[0].extract_text() or "").strip() == ""
    assert len(reader.pages[0].images) == 1


def test_rotated_scan_text_runs_across_the_short_side():
    folder, data = fixture("rotated_scan")
    image = PdfReader(folder / data["files"][0]).pages[0].images[0].image
    assert image.height > image.width  # portrait sheet
    # Ink sits in vertical bands (text lines rotated 90°): column ink profile varies far more
    # than row ink profile.
    gray = image.convert("L")
    w, h = gray.size
    px = gray.load()
    cols = [sum(px[x, y] < 128 for y in range(0, h, 4)) for x in range(0, w, 4)]
    rows = [sum(px[x, y] < 128 for x in range(0, w, 4)) for y in range(0, h, 4)]
    assert sum(c == 0 for c in cols) > 0.5 * len(cols)
    assert sum(r == 0 for r in rows) < 0.5 * len(rows)


def test_doi_paper_has_matching_doi_and_two_columns():
    folder, data = fixture("doi_paper")
    path = folder / data["files"][0]
    text = pdf_text(path)
    doi = re.search(r"10\.\d{4,9}/\S+", text).group(0)  # the 04 §2 detector regex
    assert doi == data["expected"]["metadata"]["external_ids"]["doi"]
    page = PdfReader(path).pages[0]
    width = float(page.mediabox.width)
    xs = {}

    def visit(txt, cm, tm, *_):
        if txt.strip() in ("Abstract", "References"):
            xs[txt.strip()] = cm[4] + tm[4]  # reportlab places lines through the CTM

    page.extract_text(visitor_text=visit)
    assert xs["Abstract"] < width / 2 < xs["References"]


def test_two_latex_masters():
    folder, _ = fixture("two_latex_masters")
    masters = [p.name for p in folder.glob("*.tex") if r"\documentclass" in p.read_text()]
    assert sorted(masters) == ["report_draft.tex", "report_main.tex"]
    assert r"\documentclass" not in (folder / "sections/results.tex").read_text()
    assert r"\input{sections/results}" in (folder / "report_main.tex").read_text()


def test_headerless_csv_is_numeric_from_the_first_row():
    folder, data = fixture("headerless_csv")
    rows = list(csv.reader((folder / data["files"][0]).read_text().splitlines()))
    meta = data["expected"]["metadata"]
    assert len(rows) == meta["rows"] and {len(r) for r in rows} == {meta["columns"]}
    values = [[float(c) for c in r] for r in rows]  # raises if any cell (a header) is not numeric
    times = [r[0] for r in values]
    assert all(b > a for a, b in zip(times, times[1:], strict=False))


def test_fmea_has_merged_cells_formulas_without_cache_and_two_islands():
    folder, data = fixture("merged_cells_fmea")
    path = folder / data["files"][0]
    ws = load_workbook(path).active
    assert {str(r) for r in ws.merged_cells.ranges} == {"A1:H1", "B4:B6", "B7:B9"}
    cached = load_workbook(path, data_only=True).active
    rpn = {}
    for r in range(4, 10):
        assert str(ws.cell(r, 8).value).startswith("=")
        assert cached.cell(r, 8).value is None  # forces LibreOffice recalculation (03 §C.4 step 1)
        rpn[ws.cell(r, 1).value] = ws.cell(r, 5).value * ws.cell(r, 6).value * ws.cell(r, 7).value
    assert rpn == data["expected"]["metadata"]["rpn"]
    assert all(ws.cell(r, c).value is None for r in (10, 11) for c in range(1, 9))
    assert ws["A12"].value == "Test matrix"


def test_malicious_pickle_references_exec_without_being_loaded():
    folder, data = fixture("malicious_pickle")
    raw = (folder / data["files"][0]).read_bytes()
    assert raw[:2] == b"\x80\x04"  # 03 §1.1 sniffer
    strings, globals_ = [], []
    for op, arg, _ in pickletools.genops(raw):  # parses opcodes; executes nothing
        if op.name in ("SHORT_BINUNICODE", "BINUNICODE", "UNICODE"):
            strings.append(arg)
        elif op.name == "STACK_GLOBAL":
            globals_.append(".".join(strings[-2:]))
        elif op.name == "GLOBAL":
            globals_.append(arg.replace(" ", "."))
    assert globals_ == data["expected"]["metadata"]["referenced_globals"]


def ratio(zf: zipfile.ZipFile) -> float:
    infos = zf.infolist()
    return sum(i.file_size for i in infos) / max(1, sum(i.compress_size for i in infos))


def test_zip_bomb_ratio_exceeds_guard_from_central_directory():
    folder, data = fixture("zip_bomb_ratio")
    with zipfile.ZipFile(folder / data["files"][0]) as zf:
        assert ratio(zf) > 100
        assert sum(i.file_size for i in zf.infolist()) == data["expected"]["metadata"]["uncompressed_bytes"]


def test_zip_path_traversal_member_present():
    folder, data = fixture("zip_path_traversal")
    with zipfile.ZipFile(folder / data["files"][0]) as zf:
        names = zf.namelist()
    assert any(".." in Path(n).parts or n.startswith("/") for n in names)


def test_zip_nested_depth_is_four():
    folder, data = fixture("zip_nested_depth")
    blob, depth = (folder / data["files"][0]).read_bytes(), 0
    while zipfile.is_zipfile(io.BytesIO(blob)):
        depth += 1
        with zipfile.ZipFile(io.BytesIO(blob)) as zf:
            (info,) = zf.infolist()
            assert info.file_size < 10_000  # inspect sizes before reading anything
            blob = zf.read(info)
    assert depth == data["expected"]["metadata"]["nesting_depth"]


def test_pdf_is_html_has_html_magic():
    folder, data = fixture("pdf_is_html")
    raw = (folder / data["files"][0]).read_bytes()
    assert (folder / data["files"][0]).suffix == ".pdf"
    assert not raw.startswith(b"%PDF-")
    assert raw.lstrip().lower().startswith(b"<!doctype html")


def test_unlabeled_npy_is_a_plain_float_array_without_names():
    folder, data = fixture("unlabeled_npy")
    raw = (folder / data["files"][0]).read_bytes()
    assert raw[:6] == b"\x93NUMPY"  # 03 §1.1 sniffer
    header_len = int.from_bytes(raw[8:10], "little")
    header = ast.literal_eval(raw[10 : 10 + header_len].decode("latin1"))
    meta = data["expected"]["metadata"]
    assert header == {"descr": "<f8", "fortran_order": False, "shape": tuple(meta["shape"])}
    numpy = pytest.importorskip("numpy")
    array = numpy.load(folder / data["files"][0], mmap_mode="r", allow_pickle=False)  # 03 §C.5
    assert array.dtype.names is None  # no column names: the reason for SCHEMA_UNKNOWN HARD
    assert array[0].tolist() == meta["first_row"]
    assert array.max(axis=0).tolist() == pytest.approx(meta["column_max"])

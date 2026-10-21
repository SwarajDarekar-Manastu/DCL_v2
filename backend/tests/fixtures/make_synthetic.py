"""Generate the synthetic edge-case fixtures of docs/architecture/10_development_plan.md §7.

Usage:  uv run python backend/tests/fixtures/make_synthetic.py [--out DIR]

Every fixture is fictitious (no corpus content, invented vendor, authors and numbers) and
deterministic: fixed dates, reportlab invariant mode, fixed zip timestamps. Rasterized pages
depend on the installed system fonts and are marked byte_stable: false in their expected.yaml.
The key numbers each fixture contains are the constants below; expected.yaml repeats them and
the integrity tests read them back from the generated files.
"""

import argparse
import io
import pickle
import re
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer

HERE = Path(__file__).resolve().parent
ZIP_TIME = (2026, 1, 1, 0, 0, 0)
DOC_TIME = datetime(2026, 1, 1, 0, 0, 0)
LATIN_FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
CJK_FONT = ("/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc", 0)  # index 0 = JP
PAGE_PX = (827, 1169)  # A4 at 100 dpi


# ---- helpers ---------------------------------------------------------------------------

def zip_bytes(members: list[tuple[str, bytes]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in members:
            info = zipfile.ZipInfo(name, date_time=ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            zf.writestr(info, data, compresslevel=9)
    return buf.getvalue()


def rezip_fixed_time(data: bytes) -> bytes:
    """Rewrite an OOXML package with fixed member timestamps and a fixed modified date
    (openpyxl stamps the save time into docProps/core.xml)."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        members = [(i.filename, zf.read(i.filename)) for i in zf.infolist()]
    stamp = DOC_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    members = [(name, re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + stamp, body)
                if name == "docProps/core.xml" else body) for name, body in members]
    return zip_bytes(members)


def text_image(lines: list[str], font_path, size: int, landscape: bool) -> Image.Image:
    w, h = (PAGE_PX[1], PAGE_PX[0]) if landscape else PAGE_PX
    img = Image.new("L", (w, h), 255)
    draw = ImageDraw.Draw(img)
    path, index = font_path if isinstance(font_path, tuple) else (font_path, 0)
    font = ImageFont.truetype(path, size, index=index)
    y = 80
    for line in lines:
        draw.text((80, y), line, fill=0, font=font)
        y += int(size * 1.8)
    return img


def image_only_pdf(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    c.drawImage(ImageReader(img), 0, 0, width=A4[0], height=A4[1])
    c.showPage()
    c.save()
    return buf.getvalue()


def write(out: Path, rel: str, data: bytes) -> None:
    path = out / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)


# ---- A1: rotated scan ---------------------------------------------------------------------

ROTATED_LINES = [
    "TEST CELL RUN SHEET (SYNTHETIC FIXTURE)",
    "Test Plan 07 / Run T01",
    "Chamber pressure peak: 12.4 bar",
    "Catalyst bed temperature max: 845 K",
    "Operator remarks: nominal shutdown",
]


def rotated_scan(out: Path) -> None:
    # Text laid out on a landscape sheet, then the sheet is scanned sideways: the page is
    # portrait and the text runs bottom-to-top (content rotated 90° counter-clockwise).
    img = text_image(ROTATED_LINES, LATIN_FONT, 30, landscape=True).rotate(90, expand=True)
    write(out, "a1_office/rotated_scan/run_sheet_T01_rotated.pdf", image_only_pdf(img))


# ---- A1: Japanese scanned page ------------------------------------------------------------

JAPANESE_LINES = [
    "スラスタ仕様書（合成フィクスチャ）",
    "推進剤：過酸化水素（HTP 90%）",
    "推力：1.0 N",
    "比推力：160 s",
    "燃焼室圧力：1.2 MPa",
    "質量：0.29 kg",
]


def japanese_scanned_page(out: Path) -> None:
    img = text_image(JAPANESE_LINES, CJK_FONT, 36, landscape=False)
    write(out, "a1_office/japanese_scanned_page/thruster_spec_ja_scan.pdf", image_only_pdf(img))


# ---- A1: German datasheet -----------------------------------------------------------------

GERMAN_ROWS = [
    ("Treibstoff", "HTP 90 % (Wasserstoffperoxid)"),
    ("Nennschub", "1,0 N"),
    ("Schubbereich", "0,25 - 1,1 N"),
    ("Spezifischer Impuls (Vakuum)", "160 s"),
    ("Betriebsdruck (Eintritt)", "5,5 - 22 bar"),
    ("Masse", "0,29 kg"),
    ("Ventilleistung", "8,5 W"),
    ("Temperaturbereich (Lagerung)", "-40 °C bis +85 °C"),
]


def german_datasheet(out: Path) -> None:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    c.setTitle("Datenblatt MT-1")
    c.setFont("Helvetica-Bold", 16)
    c.drawString(20 * mm, 270 * mm, "Datenblatt - Monergol-Triebwerk MT-1")
    c.setFont("Helvetica", 10)
    c.drawString(20 * mm, 262 * mm, "Beispiel Antriebstechnik GmbH (synthetische Testdatei)")
    text = c.beginText(20 * mm, 250 * mm)
    text.setFont("Helvetica", 10)
    for line in [
        "Das Triebwerk MT-1 arbeitet mit hochkonzentriertem Wasserstoffperoxid und einem",
        "Silberkatalysator. Der Nennschub von 1,0 N wird bei einem Eintrittsdruck von 22 bar",
        "erreicht. Alle Angaben sind Herstellerangaben.",
    ]:
        text.textLine(line)
    c.drawText(text)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(20 * mm, 225 * mm, "Technische Daten")
    y = 217 * mm
    for name, value in GERMAN_ROWS:
        c.setFont("Helvetica", 10)
        c.drawString(20 * mm, y, name)
        c.drawString(100 * mm, y, value)
        y -= 7 * mm
    c.setFont("Helvetica", 8)
    c.drawString(20 * mm, 15 * mm,
                 "www.beispiel-antrieb.example - Technische Änderungen vorbehalten. Ausgabe 2025-03")
    c.showPage()
    c.save()
    write(out, "a1_office/german_datasheet/MT-1_Datenblatt_DE.pdf", buf.getvalue())


# ---- A1: paper with a DOI -------------------------------------------------------------------

DOI = "10.5555/dcl.fixture.2026.001"
PAPER_REFERENCES = [
    "[1] A. Example, Catalyst bed activation in peroxide thrusters, "
    "J. Synth. Propuls. Fixtures 0 (2025) 1-9.",
    "[2] B. Sample, Start-up transients of small monopropellant thrusters, Example Press, 2024.",
    "[3] C. Placeholder, Silver screen catalysts revisited, Proc. Fixture Conf. 2023, 10-12.",
]


def doi_paper(out: Path) -> None:
    styles = getSampleStyleSheet()
    body, head = styles["BodyText"], styles["Heading2"]
    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=A4, invariant=1, title="Catalyst bed preheating (synthetic)",
                          author="A. Example; B. Sample", leftMargin=15 * mm, rightMargin=15 * mm,
                          topMargin=40 * mm, bottomMargin=15 * mm)
    gap = 6 * mm
    col_w = (doc.width - gap) / 2
    frames = [Frame(doc.leftMargin, doc.bottomMargin, col_w, doc.height, id="left"),
              Frame(doc.leftMargin + col_w + gap, doc.bottomMargin, col_w, doc.height, id="right")]

    def header(c, _doc):
        c.saveState()
        c.setFont("Helvetica", 8)
        c.drawString(15 * mm, 287 * mm, "Journal of Synthetic Propulsion Fixtures, Vol. 1 (2026) 1-4")
        c.drawString(15 * mm, 283 * mm, f"https://doi.org/{DOI}")
        c.setFont("Helvetica-Bold", 13)
        c.drawString(15 * mm, 274 * mm, "Catalyst bed preheating effects on HTP monopropellant start-up")
        c.setFont("Helvetica", 9)
        c.drawString(15 * mm, 268 * mm,
                     "A. Example, B. Sample - Example Institute of Technology (synthetic fixture)")
        c.restoreState()

    doc.addPageTemplates([PageTemplate(id="two_col", frames=frames, onPage=header)])
    story = [Paragraph("Abstract", head),
             Paragraph("We report start-up tests of a 1 N thruster fed with 90 % hydrogen peroxide. "
                       "Preheating the silver catalyst bed to 473 K shortened the pressure rise time. "
                       "The chamber pressure reached 11.8 bar at a preheat temperature of 473 K.", body),
             Paragraph("1 Introduction", head),
             Paragraph("Cold starts of peroxide thrusters show long ignition delays [1]. " * 6, body),
             Paragraph("2 Method", head),
             Paragraph("The catalyst bed was heated electrically before each firing [2]. " * 8, body),
             Paragraph("3 Results", head),
             Paragraph("The chamber pressure reached 11.8 bar at a preheat temperature of 473 K. "
                       "Without preheating the rise time was 0.8 s; "
                       "with preheating it was 0.2 s. " * 3, body),
             Paragraph("4 Conclusions", head),
             Paragraph("Preheating is recommended for repeatable start-up [3]. " * 6, body),
             Spacer(1, 4 * mm),
             Paragraph("References", head)]
    story += [Paragraph(r, body) for r in PAPER_REFERENCES]
    doc.build(story)
    write(out, "a1_office/doi_paper/preheat_paper_synthetic.pdf", buf.getvalue())


# ---- A2: two LaTeX masters ------------------------------------------------------------------

def two_latex_masters(out: Path) -> None:
    base = "a2_source/two_latex_masters/"
    write(out, base + "report_main.tex", rb"""\documentclass{article}
\usepackage{siunitx}
\title{Test Report 37 -- Hot Fire Campaign (synthetic fixture)}
\begin{document}
\maketitle
\section{Introduction}
This report summarises the hot fire campaign of Test Plan 37.
\input{sections/results}
\end{document}
""")
    write(out, base + "report_draft.tex", rb"""\documentclass{article}
\title{Test Report 37 -- DRAFT (synthetic fixture)}
\begin{document}
\maketitle
\section{Draft notes}
Early draft kept next to the final report; it is also a master document.
\end{document}
""")
    write(out, base + "sections/results.tex", rb"""\section{Results}
The chamber pressure peaked at \SI{12.4}{\bar} and the catalyst bed reached \SI{845}{\kelvin}.
""")


# ---- C: headerless CSV ------------------------------------------------------------------------

def headerless_csv_rows() -> list[list[str]]:
    rows = []
    for i in range(200):
        t = i / 1000
        rows.append([f"{t:.3f}", f"{0.5 + 0.002 * i:.4f}", f"{1.2 + 0.001 * (i % 50):.4f}",
                     f"{300 + 0.25 * i:.2f}", f"{(i % 2):d}"])
    return rows


def headerless_csv(out: Path) -> None:
    data = "\n".join(",".join(r) for r in headerless_csv_rows()) + "\n"
    write(out, "c_telemetry/headerless_csv/T05_export.csv", data.encode("ascii"))


# ---- C: merged cells (FMEA workbook) ------------------------------------------------------------

FMEA_ROWS = [  # id, subsystem (merged), failure mode, effect, S, O, D
    ("F-01", "Cooling duct", "Crack at weld", "Coolant leak", 8, 3, 4),
    ("F-02", "Cooling duct", "Blockage", "Local overheating", 7, 2, 5),
    ("F-03", "Cooling duct", "Erosion", "Wall thinning", 6, 3, 6),
    ("F-04", "Injector", "Orifice clogging", "Thrust loss", 7, 4, 3),
    ("F-05", "Injector", "Seal leak", "Propellant loss", 9, 2, 4),
    ("F-06", "Injector", "Thermal soak-back", "Valve overheating", 6, 3, 5),
]
TEST_MATRIX = [("T01", 473, 10), ("T02", 523, 10), ("T03", 573, 20)]


def merged_cells_fmea(out: Path) -> None:
    wb = Workbook()
    wb.properties.created = DOC_TIME
    wb.properties.modified = DOC_TIME
    wb.properties.creator = "synthetic fixture"
    ws = wb.active
    ws.title = "FMEA"
    ws["A1"] = "Cooling Duct FMEA (synthetic fixture)"
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:H1")
    header = ["ID", "Subsystem", "Failure mode", "Effect", "S", "O", "D", "RPN"]
    for col, name in enumerate(header, start=1):
        ws.cell(row=3, column=col, value=name).font = Font(bold=True)
    for r, (fid, sub, mode, effect, s, o, d) in enumerate(FMEA_ROWS, start=4):
        ws.cell(row=r, column=1, value=fid)
        if r in (4, 7):
            ws.cell(row=r, column=2, value=sub)
        ws.cell(row=r, column=3, value=mode)
        ws.cell(row=r, column=4, value=effect)
        ws.cell(row=r, column=5, value=s)
        ws.cell(row=r, column=6, value=o)
        ws.cell(row=r, column=7, value=d)
        ws.cell(row=r, column=8, value=f"=E{r}*F{r}*G{r}")  # no cached value: forces recalculation
    ws.merge_cells("B4:B6")
    ws.merge_cells("B7:B9")
    # second island, separated by two blank rows (rows 10-11)
    ws["A12"] = "Test matrix"
    ws["A12"].font = Font(bold=True)
    for col, name in enumerate(["Run", "Preheat (K)", "Duration (s)"], start=1):
        ws.cell(row=13, column=col, value=name).font = Font(bold=True)
    for r, row in enumerate(TEST_MATRIX, start=14):
        for col, value in enumerate(row, start=1):
            ws.cell(row=r, column=col, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    write(out, "c_telemetry/merged_cells_fmea/Cooling_Duct_FMEA.xlsx", rezip_fixed_time(buf.getvalue()))


# ---- C: malicious pickle -------------------------------------------------------------------------

PICKLE_PAYLOAD = "raise SystemExit('DCL_MALICIOUS_PICKLE_EXECUTED')"


class _Exploit:
    """Pickles to a call of builtins.exec. Harmless if ever executed (it only raises
    SystemExit with a sentinel), but any pipeline that loads it fails loudly."""

    def __reduce__(self):
        return (exec, (PICKLE_PAYLOAD,))


def malicious_pickle(out: Path) -> None:
    write(out, "c_telemetry/malicious_pickle/T03.plotdata", pickle.dumps(_Exploit(), protocol=4))


# ---- G: archive guards ----------------------------------------------------------------------------

BOMB_SIZE = 10 * 1024 * 1024


def zip_bomb_ratio(out: Path) -> None:
    write(out, "g_archives/zip_bomb_ratio/telemetry_backup.zip",
          zip_bytes([("zeros.bin", b"\0" * BOMB_SIZE)]))


def zip_path_traversal(out: Path) -> None:
    write(out, "g_archives/zip_path_traversal/scripts_bundle.zip",
          zip_bytes([("readme.txt", b"synthetic fixture\n"),
                     ("../../outside_the_scratch_dir.txt", b"must never be written\n")]))


def zip_nested_depth(out: Path) -> None:
    inner = zip_bytes([("note.txt", b"depth 4 member\n")])
    for level in (4, 3, 2):
        inner = zip_bytes([(f"level{level}.zip", inner)])
    write(out, "g_archives/zip_nested_depth/nested_archive.zip", inner)


# ---- dispatcher: .pdf that is HTML -----------------------------------------------------------------

def pdf_is_html(out: Path) -> None:
    write(out, "dispatcher/pdf_is_html/Test_Report_12.pdf", b"""<!DOCTYPE html>
<html lang="en">
<head><meta charset="utf-8"><title>Sign in required</title></head>
<body>
<h1>Sign in required</h1>
<p>This page was saved instead of the document (synthetic fixture).</p>
</body>
</html>
""")


GENERATORS = [rotated_scan, japanese_scanned_page, german_datasheet, doi_paper, two_latex_masters,
              headerless_csv, merged_cells_fmea, malicious_pickle, zip_bomb_ratio, zip_path_traversal,
              zip_nested_depth, pdf_is_html]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=HERE)
    out = parser.parse_args().out
    for generate in GENERATORS:
        generate(out)


if __name__ == "__main__":
    main()

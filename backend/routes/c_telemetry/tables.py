"""Tables of Route C (file 03 §C.3 "table", §C.4 steps 4–8): an HTML ``table`` node with ``table_row``
children that repeat the header (``Requirement R-12 | Cooling duct shall… | Verification: test``), so one
row is retrievable on its own.

Units (plan step B): the shared reader (``backend/common/unit_reader.py``) says what each column measures
and finds its unit in the header, a ``Unit`` column, a units row under the header, a row label that states
one (``Thrust [N]`` in a Parameter column), or the same column name in other tables. A numeric column is a
SOFT ``UNIT_AMBIGUITY`` only when it measures a physical quantity, or the name says nothing, and no unit
reaches the threshold; the best guess goes into the flag. Ratios, ratings, dates and IDs take no unit.
"""

import hashlib
import html
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from openpyxl.utils import get_column_letter

from backend.common.unit_reader import ColumnKind, UnitReading, name_of, read_unit
from backend.dispatcher.contract import (
    CandidateUnit,
    ChannelStats,
    Flag,
    FlagAnchor,
    ProcessedNode,
    UnitTelemetryEvidence,
)
from backend.postgres.enums import FlagSeverity, FlagType, NodeSource, NodeType

UNIT_COLUMN = re.compile(r"^units?$", re.IGNORECASE)
NAMED = re.compile(r"[^\W\d_]{2}")  # a label names something: two letters in a row, outside brackets
MAX_ROW_NODES = 2000


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.isoformat(sep=" ", timespec="seconds")
            if value.time() != time(0)
            else value.date().isoformat()
        )
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:.10g}"
    return str(value).strip()


def is_number(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, int | float):
        return True
    if isinstance(value, str):
        try:
            float(value.replace(",", "") if value.count(",") and "." in value else value.replace(",", "."))
        except ValueError:
            return False
        return bool(value.strip())
    return False


def _quantity_unit(label: str) -> str | None:
    """The unit of a physical quantity the label names (``Thrust [N]``); not a percentage or a bare value."""
    reading = read_unit(label)
    return reading.confident_unit if reading.kind == ColumnKind.QUANTITY else None


def header_unit(header: str) -> str | None:
    """The unit a header states, when the shared reader is sure of it (``Pressure (Bar) [Gauge]``)."""
    return read_unit(header).confident_unit


@dataclass
class Table:
    rows: list[list[Any]]  # first row is the header
    caption: str | None = None
    sheet: str | None = None
    cell_range: str | None = None
    hidden: bool = False
    page_note: str | None = None
    error_cells: list[str] = field(default_factory=list)
    # Where each row of ``rows`` came from, for stable anchors (file 04 §4.4): sheet row numbers with
    # the island's column span (workbooks), or (first, last) file lines (CSV; a quoted cell may wrap).
    row_lines: list[tuple[int, int]] | None = None
    col_span: tuple[int, int] | None = None

    def row_locator(self, index: int) -> dict[str, Any]:
        """Locator of ``rows[index]``: a one-row ``cell_range`` on a sheet, else the file lines."""
        if not self.row_lines or index >= len(self.row_lines):
            return {}
        first, last = self.row_lines[index]
        if self.col_span is not None:
            c0, c1 = self.col_span
            return {"cell_range": f"{get_column_letter(c0)}{first}:{get_column_letter(c1)}{last}"}
        return {"line_start": first, "line_end": last}

    def table_locator(self) -> dict[str, Any]:
        if self.col_span is not None or not self.row_lines:
            return {}
        return {"line_start": self.row_lines[0][0], "line_end": self.row_lines[-1][1]}

    @property
    def header(self) -> list[str]:
        return [cell_text(v) or f"column {i + 1}" for i, v in enumerate(self.rows[0])] if self.rows else []

    @property
    def body(self) -> list[list[Any]]:
        return self.rows[1:]

    def readings(self, same_name: Mapping[str, Sequence[str]] | None = None) -> dict[int, UnitReading]:
        """What each column measures and its unit, by the shared reader."""
        return {i: read_unit(h, same_name=same_name) for i, h in enumerate(self.header)}

    def row_units(self) -> dict[int, str]:
        """``{row index in body: unit}`` from the first text column whose labels mostly state a unit
        (``Parameter: Thrust [N]``): such a table says per row what each number measures."""
        for j in range(len(self.header)):
            labels = {
                i: cell_text(row[j]) for i, row in enumerate(self.body) if j < len(row) and cell_text(row[j])
            }
            if not labels or any(is_number(t) for t in labels.values()):
                continue
            units = {i: u for i, t in labels.items() if NAMED.search(name_of(t)) and (u := _quantity_unit(t))}
            if len(units) * 2 >= len(labels):
                return units
        return {}

    def column_units(self) -> dict[int, str]:
        units = {i: u for i, h in enumerate(self.header) if (u := header_unit(h))}
        if self.body and not self.has_unit_column():
            # A units row holds no numbers; a table with a Unit column states units per row instead.
            row = [cell_text(v) for v in self.body[0]]
            if any(t and is_number(t) for t in row):
                return units
            parsed = {}
            for i, text in enumerate(row):
                if text and not is_number(text):
                    # read with the column's name, so "C" under a temperature is degC and not coulomb
                    name = self.header[i] if i < len(self.header) else ""
                    if unit := read_unit(f"{name} [{text}]").confident_unit:
                        parsed[i] = unit
            if parsed and len(parsed) >= max(1, len([t for t in row if t]) // 2):
                units.update({i: u for i, u in parsed.items() if i not in units})
                self.rows = [self.rows[0], *self.body[1:]]  # the units row is not data
                if self.row_lines:
                    self.row_lines = [self.row_lines[0], *self.row_lines[2:]]
        return units

    def has_unit_column(self) -> bool:
        return any(UNIT_COLUMN.match(h.strip()) for h in self.header)


def table_html(table: Table) -> str:
    head = "".join(f"<th>{html.escape(h)}</th>" for h in table.header)
    body = "".join(
        "<tr>" + "".join(f"<td>{html.escape(cell_text(v))}</td>" for v in row) + "</tr>" for row in table.body
    )
    caption = f"<caption>{html.escape(table.caption)}</caption>" if table.caption else ""
    return f"<table>{caption}<thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def table_nodes(table: Table, ref: str, ordinal: int, breadcrumb: str) -> list[ProcessedNode]:
    """One ``table`` node and its ``table_row`` children (header repeated in every row)."""
    header = table.header
    units = table.column_units()
    title = " — ".join(p for p in (table.sheet, table.caption) if p)
    nodes = [
        ProcessedNode(
            ref=ref,
            node_type=NodeType.TABLE,
            ordinal=ordinal,
            breadcrumb=breadcrumb,
            content=table_html(table),
            sheet_name=table.sheet,
            cell_range=table.cell_range,
            **table.table_locator(),
            source=NodeSource.STRUCTURED,
            metadata={
                k: v
                for k, v in {
                    "caption": table.caption,
                    "hidden": table.hidden or None,
                    "units": {header[i]: u for i, u in units.items()} or None,
                    "unit_basis": {
                        header[i]: f"{r.basis} {r.confidence:.2f}"
                        for i, r in table.readings().items()
                        if i in units and r.confident_unit
                    }
                    or None,
                    "row_units": {
                        cell_text(table.body[i][0]): u for i, u in table.row_units().items() if table.body[i]
                    }
                    or None,
                    "error_cells": table.error_cells or None,
                }.items()
                if v is not None
            },
        )
    ]
    for i, row in enumerate(table.body[:MAX_ROW_NODES]):
        cells = [
            f"{header[j] if j < len(header) else f'column {j + 1}'}: {cell_text(v)}"
            for j, v in enumerate(row)
            if cell_text(v)
        ]
        if not cells:
            continue
        nodes.append(
            ProcessedNode(
                ref=f"{ref}.r{i + 1}",
                parent_ref=ref,
                node_type=NodeType.TABLE_ROW,
                level=1,
                ordinal=ordinal + 1 + i,
                breadcrumb=breadcrumb,
                content=(f"{title} | " if title else "") + " | ".join(cells),
                sheet_name=table.sheet,
                **table.row_locator(i + 1),
                source=NodeSource.STRUCTURED,
            )
        )
    return nodes


def unit_flags(table: Table, same_name: Mapping[str, Sequence[str]] | None = None) -> list[Flag]:
    """SOFT ``UNIT_AMBIGUITY`` for a numeric column that measures a physical quantity, or whose name says
    nothing, and has no unit the reader is sure of; none when the rows state their own units."""
    units = table.column_units()  # first: it drops a units row from the body
    if not table.body or table.has_unit_column() or table.row_units():
        return []
    header = table.header
    readings = table.readings(same_name)
    flags = []
    for j, name in enumerate(header):
        values = [row[j] for row in table.body if j < len(row) and cell_text(row[j])]
        if not values or j in units or readings[j].settled:
            continue
        numbers = [float(v) for v in values if isinstance(v, int | float) and not isinstance(v, bool)]
        if len(numbers) < max(1, int(0.8 * len(values))):
            continue
        reading = readings[j]
        anchor_key = hashlib.sha256(f"{table.sheet}|{table.cell_range}|{name}".encode()).hexdigest()[:16]
        flags.append(
            Flag(
                flag_type=FlagType.UNIT_AMBIGUITY,
                severity=FlagSeverity.SOFT,
                anchor=FlagAnchor(sheet=table.sheet, column=name, content_sha256=anchor_key),
                evidence=UnitTelemetryEvidence(
                    channel=name,
                    quantity_guess=str(reading.quantity) if reading.quantity else None,
                    header_lines=[" | ".join(header)],
                    stats=ChannelStats(
                        n=len(numbers), min=min(numbers), max=max(numbers), mean=sum(numbers) / len(numbers)
                    )
                    if numbers
                    else ChannelStats(),
                    candidate_units=[
                        CandidateUnit(unit=reading.unit, score=reading.confidence, reason=reading.explained)
                    ]
                    if reading.unit
                    else [],
                    note=f"numeric column {name!r} in {table.sheet or 'the table'} states no unit"
                    + (
                        f"; it reads as {reading.quantity}, usually {reading.unit}"
                        if reading.unit
                        else "; its name does not say what it measures (it may need no unit)"
                    )
                    + (f" ({reading.explained})" if reading.read_as or reading.unread else ""),
                ),
            )
        )
    return flags


def workbook_summary(tables: Sequence[Table]) -> str | None:
    """A workbook's summary: its tables by sheet and caption, with their header cells."""
    parts = []
    for table in tables[:12]:
        where = " / ".join(x for x in (table.sheet, table.caption) if x) or "table"
        header = ", ".join(c for c in (cell_text(v) for v in (table.rows[0] if table.rows else [])) if c)
        parts.append(f"{where} ({len(table.rows)} rows" + (f": {header[:160]})" if header else ")"))
    return f"Workbook with {len(tables)} tables: " + "; ".join(parts) if parts else None


def table_document(
    blob_sha256: str,
    name: str,
    context,
    tables: Sequence[Table],
    flags: Sequence[Flag],
    doc_kind,
    cir_put,
    parser_version: int,
    catalog: dict[str, Any] | None = None,
):
    """A Route C document made of tables (CSV below the series threshold, workbook islands)."""
    import json
    from pathlib import PurePosixPath

    from backend.dispatcher.contract import DocumentFields, ProcessedDocument
    from backend.postgres.enums import Route

    prefix = f"{blob_sha256}/{Route.C}/v{parser_version}"
    breadcrumb = " › ".join(context.breadcrumb)
    nodes: list[ProcessedNode] = []
    for i, table in enumerate(tables):
        nodes += table_nodes(table, f"t{i + 1}", len(nodes), breadcrumb)
    cir_put(
        f"{prefix}/tables.json",
        json.dumps(
            [
                {
                    "sheet": t.sheet,
                    "range": t.cell_range,
                    "caption": t.caption,
                    "hidden": t.hidden,
                    "rows": [[cell_text(v) for v in row] for row in t.rows],
                }  # fmt: skip
                for t in tables
            ],
            indent=1,
        ).encode(),
    )
    return ProcessedDocument(
        blob_sha256=blob_sha256,
        route=Route.C,
        parser_version=parser_version,
        cir_uri=f"s3://dcl-cir/{prefix}/",
        document=DocumentFields(
            title=PurePosixPath(name).stem,
            doc_kind=doc_kind,
            source_origin=context.source_origin,
            trust_tier=context.trust_tier,
            test_plan_id=context.test_plan_id,
            test_run_id=context.test_run_id,
            hardware_item=context.hardware_item,
            unit_count=len({t.sheet for t in tables if t.sheet}) or len(tables),
            catalog=catalog,
            summary=workbook_summary(tables),
        ),
        nodes=nodes,
        flags=list(flags),
    )
